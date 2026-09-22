#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Gerador da PLANILHA PADRÃO PS · IMPORTAÇÃO DE CADASTROS (Clientes e Fornecedores).
Fonte única do layout: qualquer mudança no modelo é mudança AQUI. Roda com openpyxl e grava
public/modelos/MODELO_importacao_cadastros_PS.xlsx (+ duas versões: _clientes e _fornecedores,
com o Tipo já sugerido). Segue o mesmo padrão aprovado da planilha de ESTOQUE.

    python3 scripts/modelos/gerar_modelo_cadastros_ps.py

Abas: Cadastros (dados), Instruções, Listas, Conferência.
As chaves técnicas (linha 4 da aba Cadastros) são o CONTRATO com o importador
(fn_cadastro_importar_previa/_aplicar e a tela /dashboard/cadastros/{clientes,fornecedores}):
há prova (prova_modelo_cadastros_ps.py) que abre este xlsx e compara.
"""
import os
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

# ── Cores PS ──────────────────────────────────────────────────────────────────
ESPRESSO = "3D2314"; DOURADO = "C8941A"; OFFWHITE = "FAF7F2"
AMARELO = "FFF4CC"; AZUL = "0000FF"; CINZA = "9C8E80"
FONT = "Arial"
DATA_INI, DATA_FIM = 5, 5005

# chave | rótulo | obrigatória | número_fmt | é_texto | descrição | no_sistema_antigo | escopo
COLS = [
    ("tipo", "Tipo", True, None, False,
     "Cliente, Fornecedor ou Ambos.", "se o cadastro é cliente, fornecedor ou os dois.", "ambos"),
    ("tipo_pessoa", "Tipo de pessoa", True, None, False,
     "PF (pessoa física) ou PJ (pessoa jurídica).", "PF/PJ (ou deduz pelo documento).", "ambos"),
    ("nome_fantasia", "Nome / Nome fantasia", True, None, False,
     "Nome (PF) ou nome fantasia (PJ).", "o nome / nome fantasia.", "ambos"),
    ("razao_social", "Razão social", False, None, False,
     "Razão social (PJ).", "a razão social.", "ambos"),
    ("cpf_cnpj", "CPF/CNPJ", False, None, True,
     "CPF (11 díg.) ou CNPJ (14). Recomendado; sem ele entra COM ALERTA. Texto, preserva zeros.",
     "o CPF ou CNPJ.", "ambos"),
    ("codigo_sistema_anterior", "Código no sistema anterior", False, None, True,
     "Código/ID do cadastro no sistema de origem (vira a referência externa p/ reimportar sem duplicar).",
     "o código do cadastro no sistema antigo.", "ambos"),
    ("sistema_origem", "Sistema de origem", False, None, False,
     "Nome do sistema de onde exportou (Omie, Conta Azul, Bling...).", "de qual sistema exportou.", "ambos"),
    ("ie", "Inscrição estadual (IE)", False, None, True,
     "IE do contribuinte, ou ISENTO. Texto.", "a inscrição estadual.", "ambos"),
    ("im", "Inscrição municipal (IM)", False, None, True,
     "Inscrição municipal, quando houver. Texto.", "a inscrição municipal.", "ambos"),
    ("contribuinte_icms", "Contribuinte ICMS", False, None, True,
     "1 contribuinte · 2 isento · 9 não contribuinte. Ver aba Listas.", "o indicador de IE do destinatário.", "ambos"),
    ("rg", "RG", False, None, True,
     "RG (pessoa física). Texto.", "o RG, se houver.", "ambos"),
    ("data_nasc_abertura", "Data de nascimento/abertura", False, "dd/mm/yyyy", False,
     "Nascimento (PF) ou data de abertura (PJ).", "a data de nascimento/abertura.", "ambos"),
    ("email", "E-mail", False, None, False,
     "E-mail principal.", "o e-mail.", "ambos"),
    ("telefone", "Telefone", False, None, True,
     "Telefone fixo. Texto.", "o telefone.", "ambos"),
    ("celular", "Celular / WhatsApp", False, None, True,
     "Celular/WhatsApp. Texto.", "o celular/WhatsApp.", "ambos"),
    ("site", "Site", False, None, False,
     "Site, quando houver.", "o site.", "ambos"),
    ("cep", "CEP", False, None, True,
     "CEP (8 díg.). Texto, preserva zeros. Preenche o endereço se faltar.", "o CEP.", "ambos"),
    ("logradouro", "Logradouro", False, None, False,
     "Rua/avenida.", "o logradouro.", "ambos"),
    ("numero", "Número", False, None, True,
     "Número (ou S/N). Texto.", "o número do endereço.", "ambos"),
    ("complemento", "Complemento", False, None, False,
     "Sala, bloco, apto...", "o complemento.", "ambos"),
    ("bairro", "Bairro", False, None, False,
     "Bairro.", "o bairro.", "ambos"),
    ("cidade", "Cidade", False, None, False,
     "Município.", "a cidade.", "ambos"),
    ("uf", "UF", False, None, True,
     "Sigla do estado (2 letras). Ver aba Listas.", "a UF.", "ambos"),
    ("pais", "País", False, None, False,
     "Padrão Brasil.", "o país (padrão Brasil).", "ambos"),
    ("codigo_ibge", "Código IBGE do município", False, None, True,
     "Opcional; é completado pelo CEP quando vazio.", "o código IBGE, se houver.", "ambos"),
    ("limite_credito", "Limite de crédito (R$)", False, "#,##0.00", False,
     "Só cliente.", "o limite de crédito.", "cliente"),
    ("condicao_pagamento", "Condição de pagamento padrão", False, None, False,
     "Ex.: à vista, 30/60/90 dias.", "a condição de pagamento padrão.", "ambos"),
    ("vendedor", "Vendedor", False, None, False,
     "Nome do vendedor responsável. Só cliente.", "o vendedor responsável.", "cliente"),
    ("segmento", "Segmento", False, None, False,
     "Segmento/ramo. Só cliente.", "o segmento do cliente.", "cliente"),
    ("categoria", "Categoria", False, None, False,
     "Categoria/grupo do cadastro.", "a categoria/grupo.", "ambos"),
    ("tags", "Tags", False, None, False,
     "Etiquetas separadas por vírgula. Só cliente.", "as tags/etiquetas.", "cliente"),
    ("banco", "Banco", False, None, False,
     "Banco do fornecedor. Só fornecedor.", "o banco do fornecedor.", "fornecedor"),
    ("agencia", "Agência", False, None, True,
     "Agência bancária. Só fornecedor. Texto.", "a agência.", "fornecedor"),
    ("conta", "Conta", False, None, True,
     "Conta bancária. Só fornecedor. Texto.", "a conta.", "fornecedor"),
    ("pix", "PIX", False, None, False,
     "Chave PIX. Só fornecedor.", "a chave PIX.", "fornecedor"),
    ("prazo_entrega_dias", "Prazo de entrega (dias)", False, "#,##0", False,
     "Prazo de entrega em dias. Só fornecedor.", "o prazo de entrega.", "fornecedor"),
    ("ativo", "Ativo", False, None, False,
     "Sim/Não (padrão Sim).", "se o cadastro está ativo.", "ambos"),
    ("observacoes", "Observações", False, None, False,
     "Anotações livres.", "as observações.", "ambos"),
]
NCOL = len(COLS)
CHAVES = [c[0] for c in COLS]

TIPOS = ["Cliente", "Fornecedor", "Ambos"]
PESSOA = ["PF", "PJ"]
SIMNAO = ["Sim", "Não"]
CONTRIB = [("1", "Contribuinte de ICMS"), ("2", "Contribuinte isento de IE"), ("9", "Não contribuinte")]
UFS = ["AC","AL","AP","AM","BA","CE","DF","ES","GO","MA","MT","MS","MG","PA","PB","PR",
       "PE","PI","RJ","RN","RS","RO","RR","SC","SP","SE","TO"]

EXEMPLO = {
    "tipo": "Ambos", "tipo_pessoa": "PJ", "nome_fantasia": "Cliente de exemplo — APAGUE esta linha",
    "razao_social": "Empresa Exemplo LTDA", "cpf_cnpj": "11222333000181",
    "codigo_sistema_anterior": "CAD-1234", "sistema_origem": "Omie", "ie": "ISENTO", "im": "",
    "contribuinte_icms": "9", "rg": "", "data_nasc_abertura": "10/03/2015",
    "email": "contato@exemplo.com.br", "telefone": "(47) 3000-0000", "celular": "(47) 99999-0000",
    "site": "www.exemplo.com.br", "cep": "89000-000", "logradouro": "Rua Exemplo", "numero": "100",
    "complemento": "Sala 2", "bairro": "Centro", "cidade": "Blumenau", "uf": "SC", "pais": "Brasil",
    "codigo_ibge": "4202404", "limite_credito": 5000, "condicao_pagamento": "30/60 dias",
    "vendedor": "João", "segmento": "Varejo", "categoria": "A", "tags": "vip, atacado",
    "banco": "Banco do Brasil", "agencia": "1234", "conta": "56789-0", "pix": "contato@exemplo.com.br",
    "prazo_entrega_dias": 7, "ativo": "Sim", "observacoes": "Cadastro de demonstração.",
}


def _title_font(sz=13, color=ESPRESSO, bold=True, italic=False):
    return Font(name=FONT, size=sz, color=color, bold=bold, italic=italic)


def build(tipo_sugerido=None):
    wb = Workbook()
    _aba_cadastros(wb.active, tipo_sugerido)
    _aba_instrucoes(wb.create_sheet("Instruções"))
    _aba_listas(wb.create_sheet("Listas"))
    _aba_conferencia(wb.create_sheet("Conferência"))
    return wb


def _aba_cadastros(ws, tipo_sugerido):
    ws.title = "Cadastros"
    last = get_column_letter(NCOL)
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]; c.value = "PLANILHA PADRÃO PS · IMPORTAÇÃO DE CLIENTES E FORNECEDORES"
    c.font = _title_font(14, ESPRESSO); c.alignment = Alignment(horizontal="left", vertical="center")
    c.fill = PatternFill("solid", fgColor=OFFWHITE); ws.row_dimensions[1].height = 24
    ws.merge_cells(f"A2:{last}2")
    c = ws["A2"]
    c.value = ("Cole os dados A PARTIR DA LINHA 5 (apague a linha de EXEMPLO). Obrigatórios: Tipo, "
               "Tipo de pessoa e Nome (colunas douradas *). CPF/CNPJ é recomendado. Não altere as linhas 3 e 4.")
    c.font = Font(name=FONT, size=10, color=ESPRESSO)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    c.fill = PatternFill("solid", fgColor=OFFWHITE); ws.row_dimensions[2].height = 28

    dourado_fill = PatternFill("solid", fgColor=DOURADO)
    espresso_fill = PatternFill("solid", fgColor=ESPRESSO)
    thin = Side(style="thin", color="D9CBB8"); border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for i, (chave, rotulo, obrig, fmt, is_text, desc, antigo, escopo) in enumerate(COLS):
        col = get_column_letter(i + 1)
        h = ws[f"{col}3"]
        h.value = rotulo + (" *" if obrig else "")
        h.font = Font(name=FONT, size=10, bold=True, color=("3D2314" if obrig else "FFFFFF"))
        h.fill = dourado_fill if obrig else espresso_fill
        h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        h.border = border
        escopo_txt = {"ambos": "Aplica a clientes e fornecedores.", "cliente": "Só para clientes.",
                      "fornecedor": "Só para fornecedores."}[escopo]
        h.comment = Comment(f"{desc}\n{escopo_txt}\nNo sistema antigo: {antigo}", "PS Gestão")
        k = ws[f"{col}4"]
        k.value = chave
        k.font = Font(name=FONT, size=9, italic=True, color=CINZA)
        k.alignment = Alignment(horizontal="center")
        e = ws[f"{col}5"]
        e.value = (tipo_sugerido if (chave == "tipo" and tipo_sugerido) else EXEMPLO[chave])
        e.font = Font(name=FONT, size=10, italic=True, color=CINZA)
        for r in range(DATA_INI, DATA_FIM + 1):
            cell = ws[f"{col}{r}"]
            if fmt:
                cell.number_format = fmt
            elif is_text:
                cell.number_format = "@"
        ws.column_dimensions[col].width = max(12, min(32, len(rotulo) + 4))

    ws.row_dimensions[3].height = 32
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:{last}{DATA_FIM}"

    colmap = {c[0]: get_column_letter(i + 1) for i, c in enumerate(COLS)}
    faixa = lambda ch: f"{colmap[ch]}{DATA_INI}:{colmap[ch]}{DATA_FIM}"
    # validações de lista (Listas!): Tipo A4:A6, Pessoa C4:C5, Contribuinte E4:E6, Sim/Não G4:G5, UF I4:I30
    dv_tipo = DataValidation(type="list", formula1="=Listas!$A$4:$A$6", allow_blank=True)
    dv_pessoa = DataValidation(type="list", formula1="=Listas!$C$4:$C$5", allow_blank=True)
    dv_contrib = DataValidation(type="list", formula1="=Listas!$E$4:$E$6", allow_blank=True)
    dv_sim = DataValidation(type="list", formula1="=Listas!$G$4:$G$5", allow_blank=True)
    dv_uf = DataValidation(type="list", formula1="=Listas!$I$4:$I$30", allow_blank=True)
    for dv in (dv_tipo, dv_pessoa, dv_contrib, dv_sim, dv_uf):
        ws.add_data_validation(dv)
    dv_tipo.add(faixa("tipo")); dv_pessoa.add(faixa("tipo_pessoa"))
    dv_contrib.add(faixa("contribuinte_icms")); dv_sim.add(faixa("ativo")); dv_uf.add(faixa("uf"))


def _aba_instrucoes(ws):
    ws.merge_cells("A1:E1")
    c = ws["A1"]; c.value = "COMO IMPORTAR CLIENTES E FORNECEDORES — 7 passos"
    c.font = _title_font(13, ESPRESSO); c.fill = PatternFill("solid", fgColor=OFFWHITE)
    passos = [
        "1) No sistema antigo, exporte a lista de clientes e/ou fornecedores (Omie: Cadastros › Clientes/Fornecedores › Exportar; "
        "Conta Azul: Cadastros › exportar Excel; Bling: Cadastros › Contatos › Exportar).",
        "2) Cole os dados na aba \"Cadastros\" A PARTIR DA LINHA 5. Não mexa nas linhas 3 e 4.",
        "3) Apague a linha de EXEMPLO. Preencha Tipo (Cliente/Fornecedor/Ambos), Tipo de pessoa (PF/PJ) e Nome.",
        "4) CPF/CNPJ é recomendado: é ele que evita duplicar e permite reimportar. Sem ele, o cadastro entra com alerta.",
        "5) Confira na aba \"Conferência\": totais por tipo, documentos vazios e duplicados na própria planilha.",
        "6) No sistema, vá em Cadastros › Clientes (ou Fornecedores) › Importar planilha e envie o arquivo. Você verá uma PRÉVIA "
        "(criar / atualizar / erro por linha) antes de confirmar.",
        "7) Reimportar a mesma planilha NÃO duplica: o casamento é por CPF/CNPJ e depois pelo código do sistema anterior.",
    ]
    r = 3
    for p in passos:
        cell = ws.cell(row=r, column=1, value=p); cell.font = Font(name=FONT, size=11)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        cell.alignment = Alignment(wrap_text=True, vertical="center"); ws.row_dimensions[r].height = 30
        r += 1
    r += 1
    hdr = ["Coluna", "Obrigatória", "Aplica a", "O que colocar", "Onde achar no sistema antigo"]
    dfill = PatternFill("solid", fgColor=DOURADO)
    for j, t in enumerate(hdr):
        cell = ws.cell(row=r, column=j + 1, value=t)
        cell.font = Font(name=FONT, size=10, bold=True, color="3D2314"); cell.fill = dfill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    r += 1
    escopo_lbl = {"ambos": "Cliente e Fornecedor", "cliente": "Só cliente", "fornecedor": "Só fornecedor"}
    for (chave, rotulo, obrig, fmt, is_text, desc, antigo, escopo) in COLS:
        ws.cell(row=r, column=1, value=rotulo).font = Font(name=FONT, size=10)
        ws.cell(row=r, column=2, value=("Sim" if obrig else "Não")).font = Font(name=FONT, size=10)
        ws.cell(row=r, column=3, value=escopo_lbl[escopo]).font = Font(name=FONT, size=10)
        ws.cell(row=r, column=4, value=desc).font = Font(name=FONT, size=10)
        ws.cell(row=r, column=5, value=antigo).font = Font(name=FONT, size=10)
        r += 1
    for col, w in (("A", 26), ("B", 12), ("C", 20), ("D", 50), ("E", 40)):
        ws.column_dimensions[col].width = w


def _aba_listas(ws):
    ws["A3"] = "Tipo do cadastro"; ws["A3"].font = _title_font(11, ESPRESSO)
    ws["C3"] = "Tipo de pessoa"; ws["C3"].font = _title_font(11, ESPRESSO)
    ws["E3"] = "Contribuinte ICMS"; ws["E3"].font = _title_font(11, ESPRESSO)
    ws["G3"] = "Sim/Não"; ws["G3"].font = _title_font(11, ESPRESSO)
    ws["I3"] = "UF"; ws["I3"].font = _title_font(11, ESPRESSO)
    for i, t in enumerate(TIPOS):
        ws.cell(row=4 + i, column=1, value=t)
    for i, t in enumerate(PESSOA):
        ws.cell(row=4 + i, column=3, value=t).number_format = "@"
    for i, (cod, sig) in enumerate(CONTRIB):
        ws.cell(row=4 + i, column=5, value=cod).number_format = "@"
        ws.cell(row=4 + i, column=6, value=sig)
    for i, t in enumerate(SIMNAO):
        ws.cell(row=4 + i, column=7, value=t)
    for i, uf in enumerate(UFS):
        ws.cell(row=4 + i, column=9, value=uf).number_format = "@"
    for col, w in (("A", 14), ("C", 8), ("E", 8), ("F", 26), ("G", 10), ("I", 8)):
        ws.column_dimensions[col].width = w


def _aba_conferencia(ws):
    # coluna A da aba Cadastros = tipo; E = cpf_cnpj (posição fixa pelo COLS)
    T = "Cadastros!$A$5:$A$5005"; DOC = "Cadastros!$E$5:$E$5005"; NM = "Cadastros!$C$5:$C$5005"
    crit_nome = f'{NM},"<>",{NM},"<>Cliente de exemplo — APAGUE esta linha"'
    ws.merge_cells("A1:D1")
    c = ws["A1"]; c.value = "CONFERÊNCIA — compare com o seu sistema antigo antes de importar"
    c.font = _title_font(12, ESPRESSO); c.fill = PatternFill("solid", fgColor=OFFWHITE)
    ws.merge_cells("A2:D2")
    c = ws["A2"]; c.value = ('Digite os números do sistema antigo na coluna amarela. "Situação" acusa OK ou DIFERENTE. '
                             'Estes indicadores são os mesmos que a importação vai mostrar na prévia.')
    c.font = Font(name=FONT, size=10, color=ESPRESSO); c.alignment = Alignment(wrap_text=True, vertical="center")
    c.fill = PatternFill("solid", fgColor=OFFWHITE); ws.row_dimensions[2].height = 28
    hdr = ["Indicador", "Nesta planilha", "No sistema antigo (digite)", "Situação"]
    dfill = PatternFill("solid", fgColor=DOURADO)
    for j, t in enumerate(hdr):
        cell = ws.cell(row=4, column=j + 1, value=t)
        cell.font = Font(name=FONT, size=10, bold=True, color="3D2314"); cell.fill = dfill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    linhas = [
        ("Total de cadastros (com nome)", f'=COUNTIFS({crit_nome})'),
        ("Clientes (Tipo = Cliente)", f'=COUNTIF({T},"Cliente")'),
        ("Fornecedores (Tipo = Fornecedor)", f'=COUNTIF({T},"Fornecedor")'),
        ("Ambos (Tipo = Ambos)", f'=COUNTIF({T},"Ambos")'),
        ("Sem CPF/CNPJ (entram com alerta)", f'=COUNTIFS({DOC},"",{crit_nome})'),
        ("CPF/CNPJ repetidos na planilha (precisa ser 0)",
         f'=SUMPRODUCT(({DOC}<>"")*(COUNTIF({DOC},{DOC})>1))'),
    ]
    amarelo = PatternFill("solid", fgColor=AMARELO)
    for k, (rot, formula) in enumerate(linhas):
        r = 5 + k
        ws.cell(row=r, column=1, value=rot).font = Font(name=FONT, size=10)
        b = ws.cell(row=r, column=2, value=formula); b.font = Font(name=FONT, size=10, bold=True)
        cc = ws.cell(row=r, column=3)
        d = ws.cell(row=r, column=4)
        if r == 10:  # repetidos
            cc.value = "—"; cc.font = Font(name=FONT, size=10, color=CINZA); cc.alignment = Alignment(horizontal="center")
            d.value = '=IF(B10=0,"OK","CORRIGIR: há documentos repetidos")'
        else:
            cc.fill = amarelo; cc.font = Font(name=FONT, size=10, color=AZUL)
            d.value = f'=IF(C{r}="","—",IF(ABS(B{r}-C{r})<0.5,"OK","DIFERENTE"))'
        d.font = Font(name=FONT, size=10, bold=True); d.alignment = Alignment(horizontal="center")
    for col, w in (("A", 44), ("B", 16), ("C", 24), ("D", 32)):
        ws.column_dimensions[col].width = w


def main():
    here = os.path.dirname(os.path.abspath(__file__))
    base = os.path.abspath(os.path.join(here, "..", "..", "public", "modelos"))
    os.makedirs(base, exist_ok=True)
    saidas = [
        ("MODELO_importacao_cadastros_PS.xlsx", None),
        ("MODELO_importacao_cadastros_PS_clientes.xlsx", "Cliente"),
        ("MODELO_importacao_cadastros_PS_fornecedores.xlsx", "Fornecedor"),
    ]
    for nome, tipo in saidas:
        out = os.path.join(base, nome)
        build(tipo).save(out)
        print("gravado:", out)


if __name__ == "__main__":
    main()
