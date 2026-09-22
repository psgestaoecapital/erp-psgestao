#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PROVA do MODELO_importacao_cadastros_PS.xlsx. Verificação ESTRUTURAL: abre o arquivo ENTREGUE e
confere (a) as chaves técnicas da aba "Cadastros" (linha 4) na ordem exata do contrato e (b) as
fórmulas da aba "Conferência" (B5:B10, C10, D5:D10). É o que garante que a planilha calcula os
mesmos indicadores que a importação vai mostrar.

    python3 scripts/modelos/prova_modelo_cadastros_ps.py

Sai 0 só se a verificação estrutural bater 100%. (O gate de build é scripts/check-modelo-cadastros.ts,
que compara a linha 4 com src/lib/cadastros/colunasImportacao.ts.)
"""
import os
import sys
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
MODELO = os.path.abspath(os.path.join(HERE, "..", "..", "public", "modelos", "MODELO_importacao_cadastros_PS.xlsx"))

CHAVES = ["tipo","tipo_pessoa","nome_fantasia","razao_social","cpf_cnpj","codigo_sistema_anterior",
          "sistema_origem","ie","im","contribuinte_icms","rg","data_nasc_abertura","email","telefone",
          "celular","site","cep","logradouro","numero","complemento","bairro","cidade","uf","pais",
          "codigo_ibge","limite_credito","condicao_pagamento","vendedor","segmento","categoria","tags",
          "banco","agencia","conta","pix","prazo_entrega_dias","ativo","observacoes"]

T = "Cadastros!$A$5:$A$5005"; DOC = "Cadastros!$E$5:$E$5005"; NM = "Cadastros!$C$5:$C$5005"
CRIT_NOME = f'{NM},"<>",{NM},"<>Cliente de exemplo — APAGUE esta linha"'
B_ESPERADO = {
    5:  f'=COUNTIFS({CRIT_NOME})',
    6:  f'=COUNTIF({T},"Cliente")',
    7:  f'=COUNTIF({T},"Fornecedor")',
    8:  f'=COUNTIF({T},"Ambos")',
    9:  f'=COUNTIFS({DOC},"",{CRIT_NOME})',
    10: f'=SUMPRODUCT(({DOC}<>"")*(COUNTIF({DOC},{DOC})>1))',
}
D_ESPERADO = {r: f'=IF(C{r}="","—",IF(ABS(B{r}-C{r})<0.5,"OK","DIFERENTE"))' for r in range(5, 10)}
D_ESPERADO[10] = '=IF(B10=0,"OK","CORRIGIR: há documentos repetidos")'


def main():
    if not os.path.exists(MODELO):
        print("modelo não existe; rode gerar_modelo_cadastros_ps.py antes", file=sys.stderr); sys.exit(2)
    wb = load_workbook(MODELO)
    assert wb.sheetnames[0] == "Cadastros", f"1ª aba deve ser Cadastros, veio {wb.sheetnames}"
    for nome in ("Cadastros", "Instruções", "Listas", "Conferência"):
        assert nome in wb.sheetnames, f"aba faltando: {nome}"
    cad = wb["Cadastros"]
    chaves = [cad.cell(row=4, column=c).value for c in range(1, len(CHAVES) + 1)]
    assert chaves == CHAVES, f"chaves técnicas (linha 4) divergem:\n  got={chaves}\n  exp={CHAVES}"

    conf = wb["Conferência"]; falhas = []
    for r, exp in B_ESPERADO.items():
        got = conf.cell(row=r, column=2).value
        if got != exp: falhas.append(f"  B{r}\n    got={got}\n    exp={exp}")
    for r, exp in D_ESPERADO.items():
        got = conf.cell(row=r, column=4).value
        if got != exp: falhas.append(f"  D{r}\n    got={got}\n    exp={exp}")
    if conf.cell(row=10, column=3).value != "—":
        falhas.append("  C10 deveria ser '—'")
    if falhas:
        print("FÓRMULAS DA CONFERÊNCIA DIVERGEM DA SPEC:\n" + "\n".join(falhas), file=sys.stderr); sys.exit(1)

    print("(A) ESTRUTURAL OK —", len(CHAVES), "chaves técnicas (linha 4):", ",".join(chaves))
    print("(B) CONFERÊNCIA OK — B5:B10, C10 e as situações D5:D10 batem com a spec.")
    print("\nPROVA OK.")


if __name__ == "__main__":
    main()
